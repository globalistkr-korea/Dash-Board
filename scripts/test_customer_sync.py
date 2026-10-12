"""Guards that prevent malformed or truncated sources from replacing published data."""
import copy
import unittest
import tempfile
import json
from pathlib import Path
from openpyxl import Workbook
from import_customer_pnl import amount, SOURCE_ID, import_workbook
from sync_customer_pnl import canonical, validate_transition, sync, FILES

class SourceGuards(unittest.TestCase):
    def setUp(self):
        self.sample = {'years':['2025','2026'], 'actualMonths':{'2025':12,'2026':8},
            'source':{'spreadsheetId':SOURCE_ID,'importedAt':'first','rowCounts':{'2025':100,'2026':80},
                      'months':{'2025':list(range(1,13)),'2026':list(range(1,9))}},
            'customers':{'customer':{}},'warehouses':{'warehouse':{}}}
    def test_import_repeat_and_invalid_header_preserve_snapshot(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);(root/'src/data').mkdir(parents=True)
            wb=Workbook();wb.remove(wb.active)
            headers=['년','월','거점','고객명','고객명 (본사)','Biz1','Biz2','Region','Revenue','Direct Cost','Direct Profit','Indirect Cost','Gross Profit','영업이익','01. HANDLING FEE']
            for year,n in [(2025,12),(2026,8)]:
                ws=wb.create_sheet(f'Customer PNL Raw {year}');ws.append(headers)
                for month in range(1,n+1):ws.append([year,month,'QA-WH','QA-Customer','QA-Customer','CL','WH','North',1000000,600000,400000,100000,300000,200000,600000])
            source=root/'sample.xlsx';wb.save(source)
            import_workbook(source,root/'src/data')
            self.assertFalse(sync(source,root))
            self.assertFalse(sync(source,root))
            before={name:(root/'src/data'/name).read_bytes() for name in FILES}
            ws=wb['Customer PNL Raw 2026'];ws.cell(1,9,'Broken Revenue');wb.save(source)
            with self.assertRaises(ValueError):sync(source,root)
            self.assertEqual(before,{name:(root/'src/data'/name).read_bytes() for name in FILES})

    def test_same_export_timestamp_is_not_a_data_change(self):
        new=copy.deepcopy(self.sample);new['source']['importedAt']='second'
        self.assertEqual(canonical(new),canonical(self.sample))
    def test_new_month_accepted(self):
        new=copy.deepcopy(self.sample);new['actualMonths']['2026']=9;new['source']['months']['2026'].append(9)
        validate_transition(self.sample,new)
    def test_removed_history_period_rows_and_gaps_rejected(self):
        for mutation in [lambda x:x['years'].remove('2025'),lambda x:x['actualMonths'].update({'2026':7}),
                         lambda x:x['source']['rowCounts'].update({'2026':20}),
                         lambda x:x['source']['months']['2026'].remove(3)]:
            with self.subTest(mutation=mutation):
                new=copy.deepcopy(self.sample);mutation(new)
                with self.assertRaises(ValueError):validate_transition(self.sample,new)
    def test_existing_numeric_conventions_and_invalid_values(self):
        self.assertEqual(amount('1.234.567'),1234567)
        self.assertEqual(amount('(1,234)'),-1234)
        self.assertEqual(amount(None),0)
        for value in ['#VALUE!','NaN',float('inf')]:
            with self.assertRaises(ValueError):amount(value)

if __name__=='__main__':unittest.main()
