"""Import the user-selected Customer PNL workbook using header-based mappings.

Usage: python3 scripts/import_customer_pnl.py workbook.xlsx
Amounts remain million VND; source rows are never filtered by customer size.
"""
import json
import math
import re
import sys
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE_ID = '1shyoDfEFd9CFlRjD4Nme6InC0-5q6QvTcI9v5iQ2bDQ'
FIELDS = {'revenue': 'Revenue', 'directCost': 'Direct Cost',
          'directProfit': 'Direct Profit', 'grossProfit': 'Gross Profit',
          'opProfit': '영업이익', 'indirectCost': 'Indirect Cost'}
WH_MAP = {'Oriental': 'Hai Nam', 'SK': 'AJ', 'DC Yen My': 'SLS'}


def amount(value):
    if value is None or str(value).strip() in ('', '-', '–'):
        return 0.0
    if isinstance(value, (int, float)):
        result = float(value)
        if not math.isfinite(result):
            raise ValueError("Non-finite source amount")
        return result
    text = str(value).strip().replace(',', '')
    if re.fullmatch(r'-?\d{1,3}(?:\.\d{3})+', text):
        text = text.replace('.', '')
    if text.startswith('(') and text.endswith(')'):
        text = '-' + text[1:-1]
    result = float(text)
    if not math.isfinite(result):
        raise ValueError("Non-finite source amount")
    return result


def import_workbook(workbook, output_dir=None):
    output_dir = Path(output_dir) if output_dir else ROOT / 'src/data'
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    years = sorted(re.search(r'(\d{4})$', name).group(1) for name in wb.sheetnames
                   if name.startswith('Customer PNL Raw '))
    if not years or len(years) != len(set(years)):
        raise ValueError('Missing or duplicate Customer PNL Raw year sheets')
    blank = lambda: {y: [0.0] * 12 for y in years}
    def entity():
        return {'regions': Counter(), 'data': {}}
    def segment():
        return {**{f: blank() for f in FIELDS}, 'items': {},
                'coverage': {y: [False] * 12 for y in years},
                'fieldCoverage': {f: {y: [False] * 12 for y in years} for f in FIELDS},
                'itemCoverage': {}, 'sourceNames': set(),
                'costCenters': {y: [set() for _ in range(12)] for y in years}}
    def remember_input(d, year, m, source_name, row, indices, cost_cols, cc_index):
        d['coverage'][year][m] = True
        d['sourceNames'].add(source_name)
        if cc_index is not None and cc_index < len(row) and row[cc_index]:
            d['costCenters'][year][m].add(re.sub(r'^(\d+)\.0$', r'\1', str(row[cc_index]).strip()))
        for f, name in FIELDS.items():
            cell = row[indices[name]] if indices[name] < len(row) else None
            if cell is not None and str(cell).strip() not in ('', '-', '–'):
                d['fieldCoverage'][f][year][m] = True
        for col, item in cost_cols:
            cell = row[col] if col < len(row) else None
            if cell is not None and str(cell).strip() not in ('', '-', '–'):
                if item not in d['itemCoverage']:
                    d['itemCoverage'][item] = {y: [False] * 12 for y in years}
                d['itemCoverage'][item][year][m] = True

    warehouses, customers = defaultdict(entity), defaultdict(entity)
    leaves = {y: defaultdict(lambda: {f: [0.0] * 12 for f in FIELDS}) for y in years}
    totals = {y: {f: [0.0] * 12 for f in FIELDS} for y in years}
    month_sets = {y: set() for y in years}
    counts, items = Counter(), set()
    unknown_regions, unknown_biz = Counter(), Counter()
    identity_errors = Counter()
    for year in years:
        rows = iter(wb[f'Customer PNL Raw {year}'].iter_rows(values_only=True))
        header = None
        for row in rows:
            if 'Revenue' in row and 'Gross Profit' in row:
                header = list(row)
                break
        if header is None:
            raise ValueError(f'{year}: missing header')
        indices = {name: header.index(name) for name in set(FIELDS.values()) |
                   {'년', '월', '거점', '고객명', '고객명 (본사)', 'Biz1', 'Biz2', 'Region'}}
        cost_cols = [(i, str(name).strip()) for i, name in enumerate(header)
                     if name and re.match(r'^\d+\.', str(name))]
        cc_index = header.index('Cost Center') if 'Cost Center' in header else None
        items.update(name for _, name in cost_cols)
        zero_rows = []
        for row in rows:
            def get(name):
                i = indices[name]
                return row[i] if i < len(row) else None
            source_year = str(get('년') or '').replace('.0', '')
            if source_year != year:
                continue
            match = re.fullmatch(r'(\d{1,2})(?:월|\.0)?', str(get('월') or '').strip())
            if not match or not 1 <= int(match.group(1)) <= 12:
                raise ValueError(f'{year}: invalid month {get("월")}')
            m = int(match.group(1)) - 1
            values = {f: amount(get(name)) / 1_000_000 for f, name in FIELDS.items()}
            costs = {name: amount(row[i] if i < len(row) else None) / 1_000_000
                     for i, name in cost_cols}
            if not any(values.values()) and not any(costs.values()):
                if any(get(name) is not None and str(get(name)).strip() not in ('', '-', '–') for name in FIELDS.values()):
                    zero_rows.append((m, row))
                continue
            month_sets[year].add(m + 1)
            region_value = str(get('Region') or '').strip()
            region = {'North': '북부', 'South': '남부', 'VNHA': '북부', 'VNHO': '남부'}.get(region_value, '미지정')
            if region == '미지정':
                unknown_regions[region_value] += 1
            clff = str(get('Biz1') or '').strip().upper()
            biz = str(get('Biz2') or '').strip().upper()
            if clff == 'FF':
                seg = 'FF'
            elif biz in ('WH', 'VAWD', 'WM'):
                seg = '창고'
            elif biz in ('LD', 'VATD', 'TM'):
                seg = '운송'
            else:
                seg = 'CL 기타' if clff == 'CL' else '기타'
                unknown_biz[f'{clff}/{biz}'] += 1
            subtype = 'WM' if seg == '창고' else 'TM' if seg == '운송' else biz or '기타'
            original_warehouse = str(get('거점') or '미지정')
            warehouse = original_warehouse.strip()
            warehouse = WH_MAP.get(warehouse, warehouse)
            customer_value = get('고객명 (본사)')
            if not customer_value or str(customer_value).startswith('#'):
                customer_value = get('고객명')
            customer = str(customer_value or '미지정').strip()
            for table, name in ((warehouses, warehouse), (customers, customer)):
                e = table[name]
                e['regions'][region] += 1
                # Separate same-name entities by region to prevent cross-region leakage.
                key = f'{region}|{seg}'
                if key not in e['data']:
                    e['data'][key] = segment()
                d = e['data'][key]
                remember_input(d, year, m, original_warehouse if table is warehouses else str(customer_value or '미지정'), row, indices, cost_cols, cc_index)
                for f, value in values.items():
                    d[f][year][m] += value
                for item, value in costs.items():
                    if not value:
                        continue
                    if item not in d['items']:
                        d['items'][item] = blank()
                    d['items'][item][year][m] += value
            leaf = leaves[year][(clff, region, subtype)]
            for f, value in values.items():
                leaf[f][m] += value
                totals[year][f][m] += value
            for name, delta in {
                'directProfit': values['revenue'] - values['directCost'] - values['directProfit'],
                'grossProfit': values['directProfit'] - values['indirectCost'] - values['grossProfit'],
            }.items():
                if abs(delta) > 0.00001:
                    identity_errors[f'{year}:{name}'] += 1
            counts[year] += 1
        # Explicit all-zero rows within a confirmed period are still entered.
        # Future templates and new zero-only entities do not alter the snapshot grain.
        for m, row in zero_rows:
            if m + 1 not in month_sets[year]:
                continue
            def zero_get(name):
                i = indices[name]
                return row[i] if i < len(row) else None
            region = {'North': '북부', 'South': '남부', 'VNHA': '북부', 'VNHO': '남부'}.get(str(zero_get('Region') or '').strip(), '미지정')
            clff = str(zero_get('Biz1') or '').strip().upper()
            biz = str(zero_get('Biz2') or '').strip().upper()
            seg = 'FF' if clff == 'FF' else '창고' if biz in ('WH','VAWD','WM') else '운송' if biz in ('LD','VATD','TM') else 'CL 기타' if clff == 'CL' else '기타'
            original_wh = str(zero_get('거점') or '미지정')
            wh = WH_MAP.get(original_wh.strip(), original_wh.strip())
            cust = zero_get('고객명 (본사)')
            if not cust or str(cust).startswith('#'):
                cust = zero_get('고객명')
            for table, name, original in ((warehouses, wh, original_wh), (customers, str(cust or '미지정').strip(), str(cust or '미지정'))):
                d = table.get(name, {}).get('data', {}).get(f'{region}|{seg}')
                if d is not None:
                    remember_input(d, year, m, original, row, indices, cost_cols, cc_index)
    def sparse(series):
        result = {}
        for y, values in series.items():
            if not any(values):
                continue
            last = max(i for i, v in enumerate(values) if v) + 1
            result[y] = [round(v, 6) for v in values[:last]]
        return result
    def finalize(table):
        out = {}
        for name, e in table.items():
            for region in e['regions']:
                data = {key.split('|', 1)[1]: {
                            **{f: sparse(val[f]) for f in FIELDS},
                            'coverage': {y: a for y, a in val['coverage'].items() if any(a)},
                            'fieldCoverage': {f: {y: a for y, a in mp.items() if any(a)} for f, mp in val['fieldCoverage'].items()},
                            'itemCoverage': {it: {y: a for y, a in mp.items() if any(a)} for it, mp in val['itemCoverage'].items()},
                            'sourceNames': sorted(val['sourceNames']),
                            'costCenters': {y: [sorted(a) for a in months] for y, months in val['costCenters'].items() if any(months)},
                            'items': {item: sparse(series) for item, series in val['items'].items() if any(any(a) for a in series.values())}}
                        for key, val in e['data'].items()
                        if key.startswith(region + '|')}
                label = name if len(e['regions']) == 1 else f'{name} ({region})'
                segs = list(data)
                out[label] = {'region': region, 'segs': segs,
                              'clff': sorted({'CL' if s in ('창고', '운송', 'CL 기타') else 'FF' if s == 'FF' else '기타' for s in segs}),
                              'data': data}
        return out
    if any(not month_sets[y] for y in years):
        raise ValueError('One or more years contain no confirmed amounts')
    source = {'spreadsheetId': SOURCE_ID, 'url': f'https://docs.google.com/spreadsheets/d/{SOURCE_ID}/edit',
              'title': 'Integrated Customer PNL_Y26.Aug_2026.09.20',
              'importedAt': datetime.now(timezone.utc).isoformat(), 'unit': '백만동',
              'rowCounts': dict(counts), 'months': {y: sorted(s) for y, s in month_sets.items()},
              'profitDefinitions': {'directProfit': 'Direct Profit = 직접이익', 'grossProfit': 'Gross Profit = 매출이익'}}
    ops = {'unit': '백만동', 'years': years, 'actualMonths': {y: max(s) for y, s in month_sets.items()},
           'source': source, 'costItems': sorted(items), 'warehouses': finalize(warehouses), 'customers': finalize(customers)}
    # Reconcile all metrics independently for both entity grains.
    for kind in ('warehouses', 'customers'):
        for y in years:
            for f in FIELDS:
                for m in range(12):
                    actual = sum((d[f].get(y, []) + [0.0] * 12)[m] for e in ops[kind].values() for d in e['data'].values())
                    assert abs(actual - totals[y][f][m]) < 0.001, (kind, y, f, m)
    mapping = {'매출': 'revenue', '매출원가': None, '직접이익': 'directProfit', '매출이익': 'grossProfit', '영업이익': 'opProfit'}
    plan = {'years': years, 'actualMonths': ops['actualMonths'], 'unit': '백만동', 'source': source, 'hasPlan': False, 'data': {}}
    for y in years:
        plan['data'][y] = {}
        for metric, field in mapping.items():
            rows = []
            for (clff, region, subtype), vals in leaves[y].items():
                values = vals[field] if field else [r - g for r, g in zip(vals['revenue'], vals['grossProfit'])]
                rows.append({'clff': clff, 'region': region, 'subtype': subtype, 'values': values})
            plan['data'][y][metric] = {'total': [sum(r['values'][i] for r in rows) for i in range(12)], 'leaves': rows}
    quality = {'source': source, 'unknownRegions': dict(unknown_regions), 'unknownBusinesses': dict(unknown_biz),
               'identityMismatchRows': dict(identity_errors), 'totalsMillionVnd': totals,
               'reconciliation': 'All years/months/metrics reconcile to warehouse and customer totals'}
    for name, data in [('ops.json', ops), ('plan3y.json', plan), ('pnl_import_quality.json', quality)]:
        (output_dir / name).write_text(json.dumps(data, ensure_ascii=False, separators=(',', ':')), encoding='utf-8')
    wb.close()
    return {'rows': dict(counts), 'months': source['months'], 'warehouses': len(ops['warehouses']),
                      'customers': len(ops['customers']), 'unknownRegions': dict(unknown_regions),
                      'unknownBusinesses': dict(unknown_biz), 'identityMismatches': dict(identity_errors)}


if __name__ == '__main__':
    print(json.dumps(import_workbook(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None), ensure_ascii=False, indent=2))
